"""Переносит навыки и Колесо жизни из Excel-файла в seed.local.json.

    python scripts/import_excel.py "путь/к/Небесный дракон.xlsx"

Полученный файл загружается в приложении: Настройки → Импорт.
Нужен пакет openpyxl (pip install openpyxl).
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

AREA_KEYS = {
    "дух": "spirit",
    "здоровье": "health",
    "деньги": "money",
    "спорт": "sport",
    "счастье": "happiness",
}

TITAN_KEYS = {
    "силовая": "strength",
    "ловкость": "agility",
    "выносливость": "endurance",
    "взрыв вын": "burst",
    "растяжка": "flexibility",
}

# лист «Навыки»: (колонка названия, колонка значения, категория)
SKILL_COLUMNS = [(1, 2, "hard"), (5, 6, "soft"), (9, 10, "hobby")]


def to_number(raw):
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


def read_areas(ws):
    areas = []
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col).value
        if not isinstance(name, str) or name.strip().lower() not in AREA_KEYS:
            continue
        name = name.strip()
        value = to_number(ws.cell(row=2, column=col).value)
        if value is None or not 0 <= value <= 100:
            print(f"  ! сфера «{name}»: значение {value} вне шкалы 0–100, ставлю 0 — поправь в приложении")
            value = 0
        areas.append({"key": AREA_KEYS[name.lower()], "name": name.capitalize(), "value": round(value)})
    return areas


def read_titans(ws):
    titans = []
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col).value
        if not isinstance(name, str) or name.strip().lower() not in TITAN_KEYS:
            continue
        value = to_number(ws.cell(row=2, column=col).value)
        titans.append({"key": TITAN_KEYS[name.strip().lower()], "value": round(min(100, max(0, value or 0)))})
    return titans


def read_goals(ws):
    """Цели (колонка B, прогресс в G) идут на небесный уровень, задачи (колонка I) — на земной."""
    goals = []
    for row in range(2, ws.max_row + 1):
        title = ws.cell(row=row, column=2).value
        if isinstance(title, str) and title.strip():
            progress = to_number(ws.cell(row=row, column=7).value) or 0
            goals.append({"title": title.strip(), "level": 1, "progress": round(min(1, max(0, progress)) * 100)})
    for row in range(2, ws.max_row + 1):
        title = ws.cell(row=row, column=9).value
        if isinstance(title, str) and title.strip():
            goals.append({"title": title.strip(), "level": 3, "progress": 0})
    return goals


def read_credits(values, formulas):
    """Кредиты: название в A, остаток в B; исходная сумма спрятана в формуле колонки C (1-ROUND(B16/100000,2))."""
    credits = []
    in_block = False
    for row in range(1, values.max_row + 1):
        label = values.cell(row=row, column=1).value
        if isinstance(label, str) and label.strip().lower().startswith("кредиты"):
            in_block = True
            continue
        if not in_block:
            continue
        formula = formulas.cell(row=row, column=3).value
        match = re.search(r"/(\d+)", formula) if isinstance(formula, str) else None
        remaining = to_number(values.cell(row=row, column=2).value)
        if not match or remaining is None or not isinstance(label, str):
            break
        credits.append({"name": label.replace("=", "").strip(), "total": int(match.group(1)), "remaining": round(remaining)})
    return credits


def read_salaries(ws):
    """Выплаты: строки вида «5го сентября | 50000» и «20го | 20000»; берём по одной на каждое число месяца."""
    by_day = {}
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row=row, column=1).value
        amount = to_number(ws.cell(row=row, column=2).value)
        match = re.match(r"\s*(\d{1,2})го", label) if isinstance(label, str) else None
        if match and amount:
            by_day.setdefault(int(match.group(1)), round(amount))
    return [{"name": "Аванс" if day >= 15 else "Зарплата", "day": day, "amount": amount} for day, amount in sorted(by_day.items())]


def read_skills(ws):
    skills = []
    for name_col, value_col, category in SKILL_COLUMNS:
        for row in range(2, ws.max_row + 1):
            name = ws.cell(row=row, column=name_col).value
            value = to_number(ws.cell(row=row, column=value_col).value)
            if not isinstance(name, str) or not name.strip() or value is None:
                continue
            skills.append({"category": category, "name": name.strip(), "value": round(min(100, max(0, value)))})
    return skills


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    source = Path(sys.argv[1])
    target = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("seed.local.json")

    # data_only: берём посчитанные значения формул, а не сами формулы
    wb = load_workbook(source, data_only=True)
    raw = load_workbook(source)  # с формулами: из них достаём исходные суммы кредитов
    seed = {
        "areas": read_areas(wb["Колесо жизни"]),
        "skills": read_skills(wb["Навыки"]),
        "titans": read_titans(wb["Тренировки"]),
        "goals": read_goals(wb["ЦЕЛИ И ЗАДАЧИ"]),
        "credits": read_credits(wb["Материальное"], raw["Материальное"]),
        "salaries": read_salaries(wb["Материальное"]),
    }
    target.write_text(json.dumps(seed, ensure_ascii=False, indent=2), encoding="utf-8")

    by_category = {}
    for s in seed["skills"]:
        by_category[s["category"]] = by_category.get(s["category"], 0) + 1
    print(f"Сфер: {len(seed['areas'])}, навыков: {len(seed['skills'])} {by_category}, титанов: {len(seed['titans'])}, целей: {len(seed['goals'])}, кредитов: {len(seed['credits'])}, выплат: {len(seed['salaries'])}")
    print(f"Записано в {target}")


if __name__ == "__main__":
    main()
